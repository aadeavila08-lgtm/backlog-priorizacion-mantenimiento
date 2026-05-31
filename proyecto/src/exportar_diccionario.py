"""
Exporta el diccionario de datos a un documento Word formal.

Cumple con el objetivo especifico 1 del informe:
"Contar con diccionario de datos completo (nombre, tipo, unidad, rango)
para cada campo".
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

from data_generator import diccionario_datos


def _color_celda(celda, hexcolor: str) -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tc_pr.append(shd)


def construir(salida: str) -> str:
    df = diccionario_datos()
    doc = Document()

    # Estilo base
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    # Titulo
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("Diccionario de Datos")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "Modelo predictivo-prescriptivo de priorizacion del backlog "
        "de mantenimiento")
    run.italic = True
    run.font.size = Pt(11)

    doc.add_paragraph()

    # Contexto
    p = doc.add_paragraph()
    p.add_run(
        "Este documento describe los campos del dataset sintetico generado "
        "por el modulo data_generator.py. Cada registro corresponde a una "
        "orden de trabajo (WO) dentro del backlog semanal y refleja los "
        "atributos tipicamente disponibles en un CMMS, alineados con la "
        "norma ISO 14224 para datos de confiabilidad y mantenimiento."
    )

    doc.add_paragraph()

    # Tabla
    tabla = doc.add_table(rows=1, cols=len(df.columns))
    tabla.style = "Light Grid Accent 1"
    tabla.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Header
    hdr = tabla.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr[i].text = col
        _color_celda(hdr[i], "1F3A5F")
        for par in hdr[i].paragraphs:
            for run in par.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Body
    for _, fila in df.iterrows():
        celdas = tabla.add_row().cells
        for j, col in enumerate(df.columns):
            celdas[j].text = str(fila[col])
            for par in celdas[j].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(9)
                    run.font.name = "Arial"

    # Anchos de columna sugeridos
    anchos_cm = [3.4, 1.8, 1.8, 2.6, 6.0]
    for fila in tabla.rows:
        for j, c in enumerate(fila.cells):
            if j < len(anchos_cm):
                c.width = Cm(anchos_cm[j])

    doc.add_paragraph()

    # Notas
    p = doc.add_paragraph()
    run = p.add_run("Notas:")
    run.bold = True

    for nota in [
        "El campo prob_real_falla_7d corresponde al 'ground truth' del "
        "simulador y NO se usa como variable de entrada del modelo: solo "
        "permite evaluar la calidad de las probabilidades predichas.",
        "El campo falla_7d (variable objetivo) se genera por muestreo "
        "Bernoulli con la probabilidad real, lo que introduce variabilidad "
        "estocastica controlada en cada replica.",
        "Los repuestos criticos se modelan como un catalogo de N ids "
        "(default = 10) con disponibilidad on-hand y lead time independiente.",
        "La habilidad requerida se deriva de la clase de equipo, reflejando "
        "la asignacion tipica en plantas de proceso.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(nota).font.size = Pt(10)

    doc.save(salida)
    return salida


if __name__ == "__main__":
    base = os.path.dirname(__file__)
    out = os.path.abspath(os.path.join(base, "..", "docs", "diccionario_datos.docx"))
    os.makedirs(os.path.dirname(out), exist_ok=True)
    print(construir(out))
